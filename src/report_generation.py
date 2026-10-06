from pathlib import Path
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[1]
REPORT_DIR = BASE_DIR / "reports"
CHART_DIR = REPORT_DIR / "charts"
REPORT_DIR.mkdir(exist_ok=True)
CHART_DIR.mkdir(exist_ok=True)

def generate_report(df, summary):
    df = df.copy()
    df["order_date"] = pd.to_datetime(df["order_date"])
    df["month"] = df["order_date"].dt.to_period("M").astype(str)

    monthly = df.groupby("month", as_index=False)["sales"].sum()
    category = df.groupby("category", as_index=False)["sales"].sum().sort_values("sales", ascending=False)
    region = df.groupby("region", as_index=False)["sales"].sum().sort_values("sales", ascending=False)
    products = df.groupby("product", as_index=False)["sales"].sum().sort_values("sales", ascending=False).head(10)

    # Create visual summaries
    plt.figure(figsize=(9, 5))
    plt.plot(monthly["month"], monthly["sales"], marker="o")
    plt.title("Monthly Sales")
    plt.xlabel("Month")
    plt.ylabel("Sales")
    plt.xticks(rotation=45)
    plt.tight_layout()
    plt.savefig(CHART_DIR / "monthly_sales.png", dpi=150)
    plt.close()

    plt.figure(figsize=(9, 5))
    plt.bar(category["category"], category["sales"])
    plt.title("Sales by Category")
    plt.xlabel("Category")
    plt.ylabel("Sales")
    plt.xticks(rotation=25)
    plt.tight_layout()
    plt.savefig(CHART_DIR / "category_sales.png", dpi=150)
    plt.close()

    plt.figure(figsize=(9, 5))
    plt.bar(region["region"], region["sales"])
    plt.title("Sales by Region")
    plt.xlabel("Region")
    plt.ylabel("Sales")
    plt.xticks(rotation=25)
    plt.tight_layout()
    plt.savefig(CHART_DIR / "region_sales.png", dpi=150)
    plt.close()

    plt.figure(figsize=(10, 6))
    plt.barh(products["product"][::-1], products["sales"][::-1])
    plt.title("Top Products by Sales")
    plt.xlabel("Sales")
    plt.tight_layout()
    plt.savefig(CHART_DIR / "top_products.png", dpi=150)
    plt.close()

    report_path = REPORT_DIR / "sales_report.xlsx"

    # Write data sheets first
    kpis = pd.DataFrame({
        "Metric": [
            "Rows Before Cleaning", "Duplicate Rows Removed",
            "Missing Cells Before", "Rows After Cleaning",
            "Missing Cells After", "Total Sales", "Total Orders",
            "Total Quantity", "Average Order Value"
        ],
        "Value": [
            summary["rows_before"], summary["duplicate_rows_removed"],
            summary["missing_cells_before"], summary["rows_after"],
            summary["missing_cells_after"], round(df["sales"].sum(), 2),
            df["order_id"].nunique(), int(df["quantity"].sum()),
            round(df["sales"].mean(), 2)
        ]
    })

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        kpis.to_excel(writer, sheet_name="Dashboard Data", index=False)
        df.drop(columns=["month"]).to_excel(writer, sheet_name="Cleaned Data", index=False)
        monthly.to_excel(writer, sheet_name="Monthly Sales", index=False)
        category.to_excel(writer, sheet_name="Category Sales", index=False)
        region.to_excel(writer, sheet_name="Region Sales", index=False)
        products.to_excel(writer, sheet_name="Top Products", index=False)

    # Polish workbook and build visual dashboard
    wb = load_workbook(report_path)
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"

    # Theme
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    title_font = Font(size=20, bold=True, color="FFFFFF")
    section_font = Font(size=12, bold=True, color="1F4E78")
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:L2")
    ws["A1"] = "AUTOMATED SALES DATA CLEANING & REPORTING DASHBOARD"
    ws["A1"].fill = dark_fill
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:L3")
    ws["A3"] = "Task 4 | Data Cleaning, Quality Validation, Analysis & Automated Reporting"
    ws["A3"].font = Font(italic=True, color="666666")
    ws["A3"].alignment = Alignment(horizontal="center")

    # KPI cards
    cards = [
        ("A5:C7", "TOTAL SALES", f"₹{df['sales'].sum():,.2f}"),
        ("D5:F7", "TOTAL ORDERS", f"{df['order_id'].nunique():,}"),
        ("G5:I7", "ROWS CLEANED", f"{summary['rows_after']:,}"),
        ("J5:L7", "DUPLICATES REMOVED", f"{summary['duplicate_rows_removed']:,}"),
    ]
    for cell_range, label, value in cards:
        ws.merge_cells(cell_range)
        start = cell_range.split(":")[0]
        c = ws[start]
        c.value = f"{label}\n{value}"
        c.fill = light_fill
        c.font = Font(size=14, bold=True, color="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    for r in range(5, 8):
        ws.row_dimensions[r].height = 24

    ws["A9"] = "DATA QUALITY SUMMARY"
    ws["A9"].font = section_font
    ws["A10"] = "Metric"
    ws["B10"] = "Before"
    ws["C10"] = "After"
    for c in ws[10][0:3]:
        c.fill = dark_fill
        c.font = white_font
        c.alignment = Alignment(horizontal="center")
    quality = [
        ("Rows", summary["rows_before"], summary["rows_after"]),
        ("Missing Cells", summary["missing_cells_before"], summary["missing_cells_after"]),
        ("Duplicates", summary["duplicate_rows_removed"], 0),
    ]
    for i, row in enumerate(quality, start=11):
        for j, value in enumerate(row, start=1):
            ws.cell(i, j, value)
            ws.cell(i, j).border = border

    # Key insight
    ws["E9"] = "KEY INSIGHTS"
    ws["E9"].font = section_font
    insights = [
        f"• Total sales generated: ₹{df['sales'].sum():,.2f}",
        f"• Highest-sales category: {category.iloc[0]['category']}",
        f"• Highest-sales region: {region.iloc[0]['region']}",
        f"• Top product: {products.iloc[0]['product']}",
        f"• Missing cells after cleaning: {summary['missing_cells_after']}",
    ]
    for i, text in enumerate(insights, start=10):
        ws.merge_cells(start_row=i, start_column=5, end_row=i, end_column=12)
        ws.cell(i, 5, text)
        ws.cell(i, 5).alignment = Alignment(vertical="center")

    # Add charts as images to dashboard
    image_positions = [
        ("A16", CHART_DIR / "monthly_sales.png"),
        ("G16", CHART_DIR / "category_sales.png"),
        ("A33", CHART_DIR / "region_sales.png"),
        ("G33", CHART_DIR / "top_products.png"),
    ]
    for pos, path in image_positions:
        img = XLImage(str(path))
        img.width = 520
        img.height = 290
        ws.add_image(img, pos)

    # Sheet formatting
    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 15
    for row in range(16, 50):
        ws.row_dimensions[row].height = 18

    # Format data sheets
    for sheet_name in ["Dashboard Data", "Cleaned Data", "Monthly Sales", "Category Sales", "Region Sales", "Top Products"]:
        sh = wb[sheet_name]
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for cell in sh[1]:
            cell.fill = dark_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, sh.max_column + 1):
            max_len = 0
            for row in sh.iter_rows(min_col=col, max_col=col):
                max_len = max(max_len, len(str(row[0].value or "")))
            sh.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 28)

    # Number formats
    for sh_name in ["Cleaned Data", "Monthly Sales", "Category Sales", "Region Sales", "Top Products"]:
        sh = wb[sh_name]
        headers = {sh.cell(1, c).value: c for c in range(1, sh.max_column + 1)}
        for header in ["sales", "unit_price"]:
            if header in headers:
                for r in range(2, sh.max_row + 1):
                    sh.cell(r, headers[header]).number_format = '₹#,##0.00'

    # Hide technical KPI sheet
    wb["Dashboard Data"].sheet_state = "hidden"

    wb.save(report_path)

    log = f"""DATA CLEANING SUMMARY
======================
Rows before cleaning: {summary['rows_before']}
Duplicate rows removed: {summary['duplicate_rows_removed']}
Missing cells before cleaning: {summary['missing_cells_before']}
Rows after cleaning: {summary['rows_after']}
Missing cells after cleaning: {summary['missing_cells_after']}
Total sales after cleaning: {summary['total_sales']}

REPORT:
reports/sales_report.xlsx

VISUAL CHARTS:
reports/charts/monthly_sales.png
reports/charts/category_sales.png
reports/charts/region_sales.png
reports/charts/top_products.png
"""
    (REPORT_DIR / "cleaning_summary.txt").write_text(log, encoding="utf-8")
    return report_path

if __name__ == "__main__":
    from data_cleaning import clean_sales_data
    df, summary = clean_sales_data()
    print(generate_report(df, summary))
